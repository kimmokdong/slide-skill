#!/usr/bin/env python3
"""HTML에서 캡처한 슬라이드 이미지(PNG)들을 순서대로 읽어
슬라이드 화면 전체에 꽉 채운 16:9 와이드 이미지형 PPTX로 내보내는 스크립트.

의존성:
    pip install python-pptx
"""

import sys
import os
import json


def assert_qa_passed(image_dir: str) -> None:
    report_path = os.path.join(image_dir, 'qa-report.json')
    if not os.path.isfile(report_path):
        raise RuntimeError(f'QA 보고서가 없습니다. capture_png.py를 먼저 실행하세요: {report_path}')
    with open(report_path, 'r', encoding='utf-8') as report_file:
        report = json.load(report_file)
    errors = report.get('summary', {}).get('errors')
    if errors is None:
        raise RuntimeError(f'QA 보고서 형식이 올바르지 않습니다: {report_path}')
    if errors:
        raise RuntimeError(f'QA 오류 {errors}건이 남아 있어 PPTX 내보내기를 중단합니다: {report_path}')

def export_image_pptx(image_dir: str, output_pptx: str) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        print("오류: python-pptx가 설치되지 않았습니다.")
        print("설치 명령: pip install python-pptx")
        sys.exit(1)

    if not os.path.exists(image_dir):
        print(f"오류: 이미지 디렉토리를 찾을 수 없습니다: {image_dir}")
        sys.exit(1)

    assert_qa_passed(image_dir)

    # 16:9 와이드 프레젠테이션 생성
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # 빈 레이아웃

    # 이미지 파일 정렬 탐색 (slide-001.png, slide-002.png 등)
    files = sorted([
        f for f in os.listdir(image_dir) 
        if f.startswith("slide-") and (f.endswith(".png") or f.endswith(".jpg"))
    ])

    if not files:
        print(f"경고: {image_dir} 경로에서 slide-*.png 형식의 슬라이드 이미지를 찾지 못했습니다.")
        sys.exit(1)

    for f in files:
        img_path = os.path.join(image_dir, f)
        slide = prs.slides.add_slide(blank_layout)
        
        # 여백 없이 슬라이드 전체에 꽉 차게 이미지 추가
        try:
            pic = slide.shapes.add_picture(img_path, Inches(0), Inches(0), height=Inches(7.5))
            # 센터 정렬 로직
            # python-pptx는 height만 지정하면 width를 비율에 맞춰 자동 조정함
            # 슬라이드 전체 너비는 13.333인치
            if pic.width.inches < 13.33:
                pic.left = Inches((13.333 - pic.width.inches) / 2)
            print(f"[EXPORT] {f} 슬라이드 삽입 완료")
        except Exception as e:
            print(f"[ERROR] {f} 이미지 삽입 중 에러 발생: {e}")

    # 출력 디렉토리 자동 생성
    output_dir = os.path.dirname(output_pptx)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    prs.save(output_pptx)
    print(f"[SUCCESS] 이미지형 PPTX 변환 완료: {output_pptx}")
    print(f"   총 슬라이드 수: {len(files)}장")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("사용법: python export_image_pptx.py <이미지디렉토리> <출력.pptx>")
        sys.exit(1)
    
    export_image_pptx(sys.argv[1], sys.argv[2])
